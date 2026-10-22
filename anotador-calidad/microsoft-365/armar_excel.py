"""Genera "Anotador Gestión de Calidad.xlsx" para OneDrive / Microsoft 365.

La hoja Tareas es una Tabla de Excel llamada "Tareas" (la que lee Power Automate).
Las columnas Días y Alerta se calculan solas; el flujo avisa las filas con Alerta.
Uso: python3 armar_excel.py salida.xlsx [email-de-prueba]
"""
import datetime as dt
import re
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableFormula, TableStyleInfo

TIPOS = 'Relevamiento,Auditoría,Capacitación,Reunión,Documento,Indicador,Recordatorio,Otro'
COLS = ['Nro', 'Cargado', 'Tarea', 'Tipo', 'Responsable', 'Vence', 'Avisar días antes', 'Agendar', 'Estado', 'Notas',
        'Días', 'Alerta', 'Calendario']
L = {c: chr(65 + i) for i, c in enumerate(COLS)}  # letra de cada columna
F_NRO = 'ROW()-ROW(Tareas[#Headers])'
F_DIAS = 'IF([@Vence]="","",[@Vence]-TODAY())'
F_ALERTA = ('IF(OR([@Estado]="Hecho",[@Vence]=""),"",IF([@Días]<0,"🔴 VENCIDA",IF([@Días]=0,"🟠 VENCE HOY",'
            'IF([@Días]<=IF([@[Avisar días antes]]="",7,[@[Avisar días antes]]),"🟡 POR VENCER",""))))')


def celda(formula):
    # En el XML de las celdas, [@Col] y [@[Col]] se guardan como Tareas[[#This Row],[Col]].
    return '=' + re.sub(r'\[@\[?([^\[\]]+)\]?\]', r'Tareas[[#This Row],[\1]]', formula)


def armar(salida, email, hoy=dt.date.today(), filas=None):
    d = lambda n: hoy + dt.timedelta(days=n)
    azul = PatternFill('solid', fgColor='1F4E79')
    wb = Workbook()

    t = wb.active
    t.title = 'Tareas'
    t.append(COLS)
    filas = filas or [
        (celda(F_NRO), hoy, 'PRUEBA – Relevamiento estación X', 'Relevamiento', 'Ana', d(-1), None, 'Sí', 'Pendiente', 'Venció ayer'),
        (celda(F_NRO), hoy, 'PRUEBA – Auditoría interna', 'Auditoría', 'Luis', d(0), None, 'Sí', 'En curso', 'Vence hoy'),
        (celda(F_NRO), hoy, 'PRUEBA – Actualizar procedimiento', 'Documento', 'Jesús', d(2), None, 'No', 'Pendiente', 'Por vencer'),
        (celda(F_NRO), hoy, 'PRUEBA – Capacitación anual', 'Capacitación', 'Equipo', d(31), 7, 'Sí', 'Pendiente', 'Más adelante'),
        (celda(F_NRO), hoy, 'PRUEBA – Revisar indicadores', 'Indicador', '', None, None, None, 'Pendiente', 'Sin fecha'),
    ]
    filas = [(celda(F_NRO),) + tuple(f[1:]) for f in filas]
    for f in filas:
        t.append(list(f) + [celda(F_DIAS), celda(F_ALERTA), None])
    ult = len(filas) + 1
    for r in range(2, ult + 1):
        t.cell(r, 2).number_format = 'DD/MM/YYYY'
        t.cell(r, 6).number_format = 'DD/MM/YYYY'
        t.cell(r, COLS.index('Días') + 1).number_format = '0'

    tabla = Table(displayName='Tareas', ref=f'A1:{L["Calendario"]}{ult}')
    tabla.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    tabla._initialise_columns()
    for col, nombre in zip(tabla.tableColumns, COLS):
        col.name = nombre
    tabla.tableColumns[0].calculatedColumnFormula = TableFormula(attr_text=F_NRO)
    tabla.tableColumns[COLS.index('Días')].calculatedColumnFormula = TableFormula(attr_text=celda(F_DIAS)[1:])
    tabla.tableColumns[COLS.index('Alerta')].calculatedColumnFormula = TableFormula(attr_text=celda(F_ALERTA)[1:])
    t.add_table(tabla)
    for col, lista in ((L['Tipo'], TIPOS), (L['Agendar'], 'Sí,No'), (L['Estado'], 'Pendiente,En curso,Hecho')):
        dv = DataValidation(type='list', formula1=f'"{lista}"', allow_blank=True)
        t.add_data_validation(dv)
        dv.add(f'{col}2:{col}1000')
    # Columnas automáticas: grises y con validación que impide escribir a mano.
    # (La protección de hoja no sirve: impide que la tabla crezca y que el flujo escriba "Calendario".)
    gris = PatternFill('solid', fgColor='E7E6E6')
    gris_titulo = PatternFill('solid', fgColor='7F7F7F')
    for col in (L['Nro'], L['Días'], L['Alerta'], L['Calendario']):
        t[f'{col}1'].fill = gris_titulo
        t[f'{col}1'].font = Font(bold=True, color='FFFFFF')
        for r in range(2, ult + 1):  # la tabla copia el formato a las filas nuevas
            t[f'{col}{r}'].fill = gris
        bloqueo = DataValidation(type='custom', formula1='FALSE', showErrorMessage=True, showInputMessage=True,
                                 errorTitle='Columna automática', error='Esta columna se completa sola: no escribir acá.',
                                 promptTitle='Automático', prompt='Se completa sola.')
        t.add_data_validation(bloqueo)
        bloqueo.add(f'{col}2:{col}1000')
    rng = f'A2:{L["Calendario"]}1000'
    t.conditional_formatting.add(rng, FormulaRule(formula=[f'${L["Estado"]}2="Hecho"'], font=Font(color='999999', strike=True), stopIfTrue=True))
    t.conditional_formatting.add(rng, FormulaRule(formula=['AND($F2<>"",$F2<TODAY())'], fill=PatternFill('solid', bgColor='F8D7DA'), stopIfTrue=True))
    t.conditional_formatting.add(rng, FormulaRule(formula=[f'${L["Alerta"]}2<>""'], fill=PatternFill('solid', bgColor='FFF3CD')))
    for i, w in enumerate([6, 12, 42, 14, 16, 12, 10, 10, 11, 34, 7, 15, 11]):
        t.column_dimensions[chr(65 + i)].width = w
    t.freeze_panes = 'A2'

    e = wb.create_sheet('Equipo')
    e.append(['Nombre', 'Email'])
    e.append(['Jesús (prueba)', email])
    for c in e[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = azul
    e.column_dimensions['A'].width = 24
    e.column_dimensions['B'].width = 34

    g = wb.create_sheet('Cómo usar')
    g.column_dimensions['A'].width = 110
    for p in ['CÓMO USAR EL ANOTADOR', '',
              '• Cargar una tarea: escribir en la primera fila vacía debajo de la tabla (la tabla se agranda sola).',
              '• Completar: Tarea, Tipo, Responsable, Vence y, si se quiere, Avisar días antes (vacío = 7 días).',
              '• Agendar = Sí: la tarea se agrega al calendario de Teams/Outlook del equipo. No o vacío: no se agenda.',
              '• Las columnas grises (Nro, Días, Alerta y Calendario) son automáticas: Excel no deja escribir en ellas.',
              '• Para borrar una tarea: clic derecho en la fila → Eliminar → Filas de tabla (no usar la tecla Supr).',
              '• Calendario la completa el flujo con "Sí" cuando crea el evento en el calendario del equipo.',
              '• Al terminar algo, poner Estado = Hecho: deja de avisarse y queda tachado.',
              '• Los avisos por mail los manda el flujo de Power Automate todos los lunes a las 8 hs.',
              '', 'La hoja Equipo es solo de referencia: los mails a los que se avisa se cargan en el flujo.']:
        g.append([p])
    g['A1'].font = Font(bold=True, size=13, color='1F4E79')
    wb.save(salida)


if __name__ == '__main__':
    armar(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else '')
